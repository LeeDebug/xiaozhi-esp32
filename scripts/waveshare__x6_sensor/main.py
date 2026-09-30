# -*- coding: utf-8 -*-
"""
Environment X6 Sensor（微雪 SKU 34169）全 API 图形化调试工具
=============================================================
- 启动后从下拉框选择串口号（不再写死端口），支持热插拔刷新
- 覆盖文档全部 8 条指令：0x70 / 0x72 / 0x74 / 0x56 / 0x54 / 0x55 / 0x71 / 0x73
- 每条指令一个测试卡片，点击即发
- 实时展示发送帧 / 接收帧，并逐字段翻译成中文含义

启动：
    uv run main.py
"""

import argparse
import os
import queue
import sys
import threading
import time
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

try:
    import serial
    from serial.tools import list_ports
except ImportError:
    sys.exit("缺少 pyserial，请先在项目目录执行：uv sync")

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

import x6_protocol as x6

# ------------------------------------------------------------------ 主题

BG        = "#f4f6f9"   # 窗口底色
CARD_BG   = "#ffffff"   # 卡片底色
LOG_BG    = "#ffffff"
FG        = "#1f2328"   # 主文字
FG_DIM    = "#6b7280"   # 次要文字
C_TX      = "#0b64c0"   # 发送 蓝
C_RX      = "#0a7d45"   # 接收 绿
C_ERR     = "#c0392b"   # 错误 红
C_INFO    = "#6b7280"   # 信息 灰
C_OK      = "#0a7d45"
C_ACCENT  = "#2563eb"

# ---- 数据存档（xlsx）----
DATA_DIR       = "data"
DATA_FILE_FMT  = "x6_log_%s.xlsx"                     # %s = YYYY-MM-DD
DATA_TITLE_FMT = "翊燊科技 环境监测系统 日志记录 %s"   # 标题日期与文件名一致
DATA_HEADERS   = ["时间 TIME", "空气质量指数 IAQ", "空气质量等级 LEVEL",
                  "总挥发物 (ppm) TVOC", "甲醛 (ppm) HCHO", "一氧化碳 (ppm) CO",
                  "温度 (℃) TEMP", "湿度 (%RH) RH"]
DATA_WIDTHS    = [21, 17, 15, 19, 16, 18, 14, 14]      # 时间 / 中文表头列加宽
DATA_NUM_FMT   = {2: "0.0", 4: "0.0000", 5: "0.0000",
                  6: "0.0000", 7: "0.00", 8: "0.00"}
DATA_FLUSH_ROWS  = 20    # 累计多少行落盘一次
DATA_FLUSH_SECS  = 30.0  # 或隔多少秒落盘一次

FONT      = ("Microsoft YaHei UI", 9)
FONT_B    = ("Microsoft YaHei UI", 9, "bold")
FONT_H    = ("Microsoft YaHei UI", 11, "bold")
FONT_BIG  = ("Microsoft YaHei UI", 16, "bold")
FONT_MONO = ("Consolas", 9)


def now_str():
    return datetime.now().strftime("%H:%M:%S.") + "%03d" % (datetime.now().microsecond // 1000)


def list_serial_ports():
    """枚举本机串口，返回 [(device, description), ...]"""
    out = []
    try:
        for p in sorted(list_ports.comports(), key=lambda c: str(c.device)):
            desc = (p.description or "").strip()
            if desc in ("n/a", ""):
                desc = p.hwid or "串口设备"
            out.append((str(p.device), desc))
    except Exception:
        pass
    return out


# ------------------------------------------------------------------ 串口工作线程


class SerialWorker(threading.Thread):
    """独占串口的后台线程：串行处理请求，避免多发串帧。"""

    def __init__(self, port, baud, logq):
        super().__init__(daemon=True)
        self.port = port
        self.baud = baud
        self.logq = logq
        self.req_q = queue.Queue()
        self.stop_evt = threading.Event()
        self.poll_on = False
        self.poll_interval = 1.0
        self.ser = None
        self._next_poll = 0.0

    # ---------------- 对外
    def submit(self, name, tx, exp_len, expect_cmd=None):
        self.req_q.put((name, bytes(tx), int(exp_len), expect_cmd))

    def set_poll(self, on, interval):
        self.poll_on = bool(on)
        self.poll_interval = max(0.2, float(interval))
        self._next_poll = 0.0

    def stop(self):
        self.stop_evt.set()

    # ---------------- 内部
    def _emit(self, kind, title, hexs=None, lines=None):
        self.logq.put({
            "kind": kind, "ts": now_str(), "title": title,
            "hex": hexs, "lines": list(lines or []),
        })

    def run(self):
        try:
            self.ser = serial.Serial(port=self.port, baudrate=self.baud, bytesize=8,
                                     parity="N", stopbits=1, timeout=0.2,
                                     write_timeout=1.0)
        except Exception as e:
            self._emit("status", "打开串口失败：%s" % e,
                       lines=["请确认：端口是否被其他软件（SSCOM/Arduino IDE）占用、"
                              "驱动是否安装、端口号是否正确。"])
            self.logq.put({"kind": "closed"})
            return

        self.logq.put({"kind": "opened", "ts": now_str(),
                       "title": "已打开 %s @ %d 8N1" % (self.port, self.baud)})

        while not self.stop_evt.is_set():
            item = None
            try:
                item = self.req_q.get(timeout=0.05)
            except queue.Empty:
                if self.poll_on and time.time() >= self._next_poll:
                    tx, exp = x6.make_request(x6.CMD_CONCENTRATION)
                    item = ("查询当前浓度（自动轮询）", tx, exp, x6.CMD_CONCENTRATION)
                    self._next_poll = time.time() + self.poll_interval

            if item:
                self._transact(item)
            else:
                try:
                    extra = self.ser.read(64)
                except Exception:
                    break
                if extra:
                    d = x6.decode(extra)
                    self._emit("rx", "异步接收（非本次请求的数据）",
                               x6.hexstr(extra),
                               ["%s：%s" % (f[0], f[2]) for f in d["fields"]])

        try:
            if self.ser and self.ser.is_open:
                self.ser.close()
        except Exception:
            pass
        self.logq.put({"kind": "closed"})

    def _transact(self, item):
        name, tx, exp_len, expect_cmd = item
        if not (self.ser and self.ser.is_open):
            self._emit("err", "%s —— 串口未打开" % name)
            return

        self._emit("tx", "发送 · %s" % name, x6.hexstr(tx), x6.describe_tx(tx))
        try:
            self.ser.reset_input_buffer()
            self.ser.write(tx)
        except Exception as e:
            self._emit("err", "%s —— 写入失败：%s" % (name, e))
            return

        buf = bytearray()
        deadline = time.time() + 2.0
        if exp_len > 0:
            # 已知应答长度：凑够就收工
            while time.time() < deadline and len(buf) < exp_len:
                try:
                    chunk = self.ser.read(exp_len - len(buf)) or b""
                except Exception as e:
                    self._emit("err", "%s —— 读取失败：%s" % (name, e))
                    return
                if chunk:
                    buf += chunk
                if self.stop_evt.is_set():
                    return
        else:
            # 未知长度（自定义帧）：读到 0.3 秒静默为止
            last = time.time()
            while time.time() < deadline:
                try:
                    chunk = self.ser.read(64) or b""
                except Exception as e:
                    self._emit("err", "%s —— 读取失败：%s" % (name, e))
                    return
                if chunk:
                    buf += chunk
                    last = time.time()
                elif buf and time.time() - last > 0.3:
                    break
                if self.stop_evt.is_set():
                    return

        if not buf:
            self._emit("err", "接收 · %s —— 超时无响应" % name, None, [
                "2 秒内没有收到任何字节。",
                "排查：TX/RX 是否接反、供电是否为 5V、波特率是否 9600、传感器是否处于休眠（先发 0x55 唤醒）。",
            ])
            return

        rx = bytes(buf)
        d = x6.decode(rx)
        if expect_cmd is not None and rx and rx[0] != expect_cmd:
            d["fields"].insert(0, ("指令头不匹配", "0x%02X" % rx[0],
                                   "期望 0x%02X，可能收到了上一条指令的残留数据" % expect_cmd))

        tag = "校验通过" if d["checksum_ok"] else "校验失败"
        lines = ["%s：%s" % (f[0], f[2]) for f in d["fields"]]
        self._emit("rx" if d["checksum_ok"] else "err",
                   "接收 · %s（%s）" % (name, tag), x6.hexstr(rx), lines)
        self._emit("summary", d["summary"] or name, None, None)
        if d.get("values"):
            self.logq.put({"kind": "values", "values": d["values"]})


# ------------------------------------------------------------------ 可滚动容器


class ScrollFrame(ttk.Frame):
    def __init__(self, master, root):
        super().__init__(master)
        self.root = root
        self.canvas = tk.Canvas(self, highlightthickness=0, bg=BG, bd=0)
        self.vsb = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)

        self.inner.bind("<Configure>",
                        lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self._win = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.bind("<Configure>",
                         lambda e: self.canvas.itemconfig(self._win, width=e.width))
        self.canvas.configure(yscrollcommand=self.vsb.set)

        self.canvas.pack(side="left", fill="both", expand=True)
        self.vsb.pack(side="right", fill="y")
        self.root.bind_all("<MouseWheel>", self._on_wheel, add="+")

    def _on_wheel(self, event):
        w = self.root.winfo_containing(event.x_root, event.y_root)
        if w is None:
            return
        s = str(w)
        if s == str(self.canvas) or s.startswith(str(self.inner) + ".") or s == str(self.inner):
            self.canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")
            return "break"


# ------------------------------------------------------------------ 主界面


class App:
    def __init__(self, root, preset_port=None):
        self.root = root
        self.logq = queue.Queue()
        self.worker = None
        self.port_map = {}
        self._poll_var = tk.BooleanVar(value=False)
        self.current_iaq = None
        self._level_color = FG_DIM
        self._iaq_win = None
        self.wb = None            # openpyxl Workbook（按天一个文件）
        self.ws = None
        self.data_date = None     # 当前存档对应的日期
        self.data_path = None
        self._pending_rows = 0
        self._last_save = 0.0
        self._save_warned = False

        root.title("Environment X6 Sensor · 全 API 调试工具  (微雪 SKU 34169)")
        root.geometry("1280x860")
        root.configure(bg=BG)
        root.minsize(1060, 700)

        self._style()
        self._build_serial_bar(preset_port)
        self._build_body()
        self._build_values()
        self._build_status()

        self.refresh_ports(select=preset_port)
        self._after_id = root.after(80, self._drain_log)
        root.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------------------------------------------------------- 样式
    def _style(self):
        st = ttk.Style()
        try:
            st.theme_use("clam")
        except tk.TclError:
            pass
        st.configure(".", background=BG, foreground=FG, font=FONT)
        st.configure("TFrame", background=BG)
        st.configure("Card.TFrame", background=CARD_BG, relief="flat")
        st.configure("TLabelframe", background=CARD_BG, relief="solid", borderwidth=1)
        st.configure("TLabelframe.Label", background=CARD_BG, foreground=C_ACCENT,
                     font=FONT_B)
        st.configure("TLabel", background=CARD_BG, foreground=FG)
        st.configure("Head.TLabel", background=CARD_BG, foreground=FG_DIM, font=FONT)
        st.configure("Unit.TLabel", background=CARD_BG, foreground="#9aa1ac",
                     font=("Microsoft YaHei UI", 8))
        st.configure("Hint.TLabel", background=CARD_BG, foreground=C_ACCENT,
                     font=("Microsoft YaHei UI", 8, "underline"))
        st.configure("TButton", padding=(10, 5))
        st.configure("Accent.TButton", padding=(10, 5))
        st.map("Accent.TButton", background=[("active", "#1d4ed8"), ("!disabled", C_ACCENT)],
               foreground=[("!disabled", "#ffffff")])
        st.configure("TCheckbutton", background=CARD_BG)
        st.configure("Bar.TCheckbutton", background=BG, foreground=FG)
        st.configure("TRadiobutton", background=CARD_BG)
        st.configure("TCombobox", padding=(4, 3))

    # ---------------------------------------------------------- 顶部：串口
    def _build_serial_bar(self, preset_port):
        bar = ttk.Frame(self.root, padding=(10, 8))
        bar.pack(side="top", fill="x")
        card = ttk.Frame(bar, style="Card.TFrame", padding=(12, 10))
        card.pack(fill="x")

        ttk.Label(card, text="串口", font=FONT_H).grid(row=0, column=0, rowspan=2, padx=(0, 14))

        ttk.Label(card, text="端口号").grid(row=0, column=1, sticky="e", padx=(0, 6))
        self.port_var = tk.StringVar()
        self.port_cb = ttk.Combobox(card, textvariable=self.port_var, width=42,
                                    state="normal", font=FONT)
        self.port_cb.grid(row=0, column=2, sticky="w")

        ttk.Label(card, text="波特率").grid(row=0, column=3, sticky="e", padx=(16, 6))
        self.baud_var = tk.StringVar(value=str(x6.BAUDRATE))
        self.baud_cb = ttk.Combobox(card, textvariable=self.baud_var, width=10,
                                    state="readonly", font=FONT,
                                    values=["1200", "2400", "4800", "9600",
                                            "19200", "38400", "57600", "115200"])
        self.baud_cb.grid(row=0, column=4, sticky="w")

        self.open_btn = ttk.Button(card, text="打开串口", style="Accent.TButton",
                                   command=self.toggle_port)
        self.open_btn.grid(row=0, column=5, padx=(16, 6))

        ttk.Button(card, text="刷新端口", command=lambda: self.refresh_ports()).grid(
            row=0, column=6, padx=(0, 6))

        self.state_var = tk.StringVar(value="未连接")
        self.state_lbl = ttk.Label(card, textvariable=self.state_var, font=FONT_B,
                                   foreground=FG_DIM, background=CARD_BG)
        self.state_lbl.grid(row=0, column=7, padx=(10, 0))

        ttk.Label(card, text="提示：端口号在启动后从下拉框选择，支持热插拔后点「刷新端口」；"
                            "未列出时可手动输入（如 COM12）。",
                 font=FONT, foreground=FG_DIM,
                 background=CARD_BG).grid(row=1, column=1, columnspan=7, sticky="w", pady=(6, 0))

    # ---------------------------------------------------------- 中部
    def _build_body(self):
        body = ttk.Frame(self.root, padding=(10, 0))
        body.pack(side="top", fill="both", expand=True)

        # ---- 左：API 卡片
        left = ttk.Frame(body, width=540)
        left.pack(side="left", fill="both", expand=False)
        left.pack_propagate(False)
        self.scroll = ScrollFrame(left, self.root)
        self.scroll.pack(fill="both", expand=True)
        api = self.scroll.inner

        self._card_concentration(api)
        self._card_param(api)
        self._card_led_get(api)
        self._card_led_set(api)
        self._card_sleep(api)
        self._card_wakeup(api)
        self._card_sn(api)
        self._card_version(api)
        self._card_raw(api)

        # ---- 右：日志
        right = ttk.Frame(body)
        right.pack(side="left", fill="both", expand=True, padx=(10, 0))

        head = ttk.Frame(right)
        head.pack(fill="x")
        ttk.Label(head, text="实时通信日志（发送 / 接收 / 中文解析）",
                  font=FONT_H, background=BG).pack(side="left", pady=(0, 6))
        ttk.Button(head, text="清空", command=self.clear_log).pack(side="right", padx=(6, 0))
        ttk.Button(head, text="保存日志", command=self.save_log).pack(side="right")
        self.autoscroll_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(head, text="自动滚动", style="Bar.TCheckbutton",
                        variable=self.autoscroll_var).pack(side="right", padx=(0, 12))

        wrap = ttk.Frame(right, style="Card.TFrame", padding=2)
        wrap.pack(fill="both", expand=True)

        self.log = tk.Text(wrap, wrap="none", bg=LOG_BG, fg=FG, font=FONT_MONO,
                           relief="flat", bd=0, padx=10, pady=8, spacing1=2, spacing3=2,
                           insertbackground=FG)
        ysb = ttk.Scrollbar(wrap, orient="vertical", command=self.log.yview)
        xsb = ttk.Scrollbar(wrap, orient="horizontal", command=self.log.xview)
        self.log.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        self.log.grid(row=0, column=0, sticky="nsew")
        ysb.grid(row=0, column=1, sticky="ns")
        xsb.grid(row=1, column=0, sticky="ew")
        wrap.rowconfigure(0, weight=1)
        wrap.columnconfigure(0, weight=1)

        for tag, color in (("ts", "#9aa1ac"), ("tx", C_TX), ("rx", C_RX),
                           ("err", C_ERR), ("info", C_INFO), ("sum", C_ACCENT),
                           ("mono", FG), ("dim", FG_DIM), ("ok", C_OK)):
            self.log.tag_config(tag, foreground=color)
        self.log.tag_config("tx", font=("Consolas", 9, "bold"))
        self.log.tag_config("rx", font=("Consolas", 9, "bold"))
        self.log.tag_config("sum", font=("Microsoft YaHei UI", 9, "bold"))
        self.log.tag_config("ts", font=("Consolas", 8))
        self.log.configure(state="disabled")

        self.log.bind("<MouseWheel>",
                      lambda e: "break" if self._log_wheel(e) else None)

    def _log_wheel(self, event):
        self.log.yview_scroll(int(-1 * (event.delta / 120)), "units")
        return True

    # ---------------------------------------------------------- 卡片工厂
    def _card(self, parent, title, note):
        f = ttk.LabelFrame(parent, text="  " + title + "  ", padding=(10, 8))
        f.pack(fill="x", padx=4, pady=5)
        if note:
            ttk.Label(f, text=note, style="Head.TLabel", wraplength=470,
                      justify="left").pack(anchor="w", pady=(0, 6))
        row = ttk.Frame(f, style="Card.TFrame")
        row.pack(fill="x")
        return f, row

    def _card_concentration(self, p):
        f, row = self._card(p, "0x70  查询当前浓度",
                            "读取 IAQ / TVOC / HCHO / CO 实时浓度 + 温度 + 湿度（返回 22 字节）。")
        ttk.Button(row, text="查询一次", style="Accent.TButton",
                   command=lambda: self._send(x6.CMD_CONCENTRATION, "查询当前浓度")
                   ).pack(side="left")

        self._poll_var = tk.BooleanVar(value=False)
        self.poll_btn = ttk.Button(row, text="开始连续采集", command=self.toggle_poll)
        self.poll_btn.pack(side="left", padx=(8, 0))
        ttk.Label(row, text="间隔", background=CARD_BG).pack(side="left", padx=(12, 4))
        self.interval_var = tk.StringVar(value="1.0")
        ttk.Spinbox(row, from_=0.2, to=60.0, increment=0.5, width=6,
                    textvariable=self.interval_var, format="%.1f").pack(side="left")
        ttk.Label(row, text="秒", background=CARD_BG).pack(side="left", padx=(2, 0))

        self.csv_var = tk.BooleanVar(value=True)   # 默认开启记录
        ttk.Checkbutton(row, text="同时记录（data/ 目录按天存档）",
                        variable=self.csv_var,
                        command=self._toggle_csv).pack(side="left", padx=(14, 0))
        return f

    def _card_param(self, p):
        f, row = self._card(p, "0x72  读取传感器参数",
                            "查询指定通道的检测量程与浓度单位（返回 7 字节）。")
        ttk.Label(row, text="通道", background=CARD_BG).pack(side="left")
        self.gas_var = tk.StringVar(value="IAQ (0x00)")
        cb = ttk.Combobox(row, textvariable=self.gas_var, width=14, state="readonly",
                          values=["IAQ (0x00)", "TVOC (0x01)", "HCHO (0x02)", "CO (0x03)"])
        cb.pack(side="left", padx=(6, 8))
        ttk.Button(row, text="读取", style="Accent.TButton",
                   command=self._send_param).pack(side="left")
        ttk.Button(row, text="读取全部 4 个通道",
                   command=self._send_param_all).pack(side="left", padx=(8, 0))

    def _card_led_get(self, p):
        f, row = self._card(p, "0x74  读取 LED 状态",
                            "查询运行指示灯当前状态：0x00 = 关闭，其他值 = 闪烁（返回 3 字节）。")
        ttk.Button(row, text="读取 LED 状态", style="Accent.TButton",
                   command=lambda: self._send(x6.CMD_LED_GET, "读取 LED 状态")
                   ).pack(side="left")

    def _card_led_set(self, p):
        f, row = self._card(p, "0x56  设置运行灯状态",
                            "控制运行指示灯开关，成功返回 \"OK\"（返回 5 字节）。")
        self.led_var = tk.StringVar(value="开启 (0x01)")
        cb = ttk.Combobox(row, textvariable=self.led_var, width=14, state="readonly",
                          values=["关闭 (0x00)", "开启 (0x01)"])
        cb.pack(side="left")
        ttk.Button(row, text="设置", style="Accent.TButton",
                   command=self._send_led_set).pack(side="left", padx=(8, 0))

    def _card_sleep(self, p):
        f, row = self._card(p, "0x54  进入休眠模式",
                            "发送 ASCII \"sleep\" 让传感器进入低功耗休眠；休眠后不再响应 0x70，"
                            "需用 0x55 唤醒。")
        ttk.Button(row, text="发送休眠指令", style="Accent.TButton",
                   command=lambda: self._send(x6.CMD_SLEEP, "进入休眠模式", b"sleep")
                   ).pack(side="left")

    def _card_wakeup(self, p):
        f, row = self._card(p, "0x55  退出休眠模式",
                            "发送 ASCII \"awake\" 唤醒传感器，恢复正常工作。")
        ttk.Button(row, text="发送唤醒指令", style="Accent.TButton",
                   command=lambda: self._send(x6.CMD_WAKEUP, "退出休眠模式", b"awake")
                   ).pack(side="left")

    def _card_sn(self, p):
        f, row = self._card(p, "0x71  获取传感器序列号 SN",
                            "读取 6 字节唯一序列号（返回 8 字节）。")
        ttk.Button(row, text="读取 SN", style="Accent.TButton",
                   command=lambda: self._send(x6.CMD_SN, "获取传感器序列号")
                   ).pack(side="left")

    def _card_version(self, p):
        f, row = self._card(p, "0x73  获取软件版本号",
                            "读取 19 字节 ASCII 固件版本字符串（返回 21 字节）。")
        ttk.Button(row, text="读取版本号", style="Accent.TButton",
                   command=lambda: self._send(x6.CMD_VERSION, "获取软件版本号")
                   ).pack(side="left")

    def _card_raw(self, p):
        f, row = self._card(p, "自定义原始帧（调试用）",
                            "直接发送十六进制帧，自动补 0-add8 校验和；勾选后需自己带校验字节。")
        self.raw_var = tk.StringVar(value="70")
        ttk.Entry(row, textvariable=self.raw_var, width=34, font=FONT_MONO).pack(side="left")
        self.raw_chk_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(row, text="已含校验和", variable=self.raw_chk_var).pack(side="left", padx=(8, 0))
        ttk.Button(row, text="发送", command=self._send_raw).pack(side="left", padx=(8, 0))

    # ---------------------------------------------------------- 底部：数值面板
    def _build_values(self):
        box = ttk.Frame(self.root, padding=(10, 6))
        box.pack(side="top", fill="x")
        card = ttk.Frame(box, style="Card.TFrame", padding=(12, 10))
        card.pack(fill="x")

        head = ttk.Frame(card, style="Card.TFrame")
        head.grid(row=0, column=0, sticky="w", padx=(0, 14))
        ttk.Label(head, text="实时数值", font=FONT_H, background=CARD_BG).pack(anchor="w")
        ttk.Label(head, text="REAL-TIME", style="Unit.TLabel").pack(anchor="w")

        # 每格：第一行 中文 + 单位，第二行 英文，第三行 数值
        self.val_vars = {}
        specs = [
            ("iaq",  "空气质量指数",    "IAQ"),
            ("tvoc", "总挥发物 (ppm)",  "TVOC"),
            ("hcho", "甲醛 (ppm)",      "HCHO"),
            ("co",   "一氧化碳 (ppm)",  "CO"),
            ("temp", "温度 (℃)",       "TEMP"),
            ("hum",  "湿度 (%RH)",      "RH"),
        ]
        for i, (key, cn, en) in enumerate(specs, start=1):
            cell = ttk.Frame(card, style="Card.TFrame")
            cell.grid(row=0, column=i, padx=7, sticky="w")
            ttk.Label(cell, text=cn, style="Head.TLabel").pack(anchor="w")
            ttk.Label(cell, text=en, style="Unit.TLabel").pack(anchor="w")
            v = tk.StringVar(value="—")
            tk.Label(cell, textvariable=v, font=FONT_BIG, fg=FG,
                     bg=CARD_BG).pack(anchor="w")
            self.val_vars[key] = v

        # IAQ 等级（点击查看详细说明）
        lvl_cell = ttk.Frame(card, style="Card.TFrame")
        lvl_cell.grid(row=0, column=len(specs) + 1, padx=(12, 0), sticky="w")
        ttk.Label(lvl_cell, text="空气质量等级", style="Head.TLabel").pack(anchor="w")
        ttk.Label(lvl_cell, text="IAQ LEVEL", style="Unit.TLabel").pack(anchor="w")
        self.level_var = tk.StringVar(value="—")
        self.level_lbl = tk.Label(lvl_cell, textvariable=self.level_var, font=FONT_B,
                                  fg=FG_DIM, bg=CARD_BG, cursor="hand2")
        self.level_lbl.pack(anchor="w")
        self.level_lbl.bind("<Button-1>", lambda e: self.show_iaq_help())
        hint = tk.Label(lvl_cell, text="点击查看等级说明 ▾", font=("Microsoft YaHei UI", 8),
                        fg=C_ACCENT, bg=CARD_BG, cursor="hand2")
        hint.pack(anchor="w")
        for w in (self.level_lbl, hint):
            w.bind("<Enter>", lambda e, wgt=w: wgt.configure(fg=C_ACCENT))
            w.bind("<Leave>", lambda e, wgt=w: wgt.configure(
                fg=C_ACCENT if wgt is hint else self._level_color))

        self.extra_var = tk.StringVar(value="")
        ttk.Label(card, textvariable=self.extra_var, style="Head.TLabel",
                  wraplength=440, justify="left").grid(
            row=1, column=1, columnspan=len(specs) + 1, sticky="w", pady=(8, 0))

    # ---------------------------------------------------------- IAQ 等级说明
    def show_iaq_help(self):
        old = getattr(self, "_iaq_win", None)
        if old is not None and old.winfo_exists():
            old.destroy()
            self._iaq_win = None
            return

        win = tk.Toplevel(self.root)
        self._iaq_win = win
        win.title("IAQ 等级说明（EPA 标准）")
        win.configure(bg=CARD_BG)
        win.transient(self.root)
        win.resizable(False, False)

        x = max(40, self.level_lbl.winfo_rootx() - 470)
        y = self.level_lbl.winfo_rooty() + 28
        win.geometry("+%d+%d" % (x, y))

        wrap = ttk.Frame(win, style="Card.TFrame", padding=(16, 12))
        wrap.pack(fill="both", expand=True)

        ttk.Label(wrap, text="IAQ（室内空气质量指数）等级对照表",
                  font=FONT_H, background=CARD_BG).pack(anchor="w", pady=(0, 2))
        ttk.Label(wrap, text="等级划分遵循 EPA 标准，数值越低空气质量越好。",
                  style="Head.TLabel").pack(anchor="w", pady=(0, 10))

        grid = ttk.Frame(wrap, style="Card.TFrame")
        grid.pack(fill="x")
        heads = [("IAQ 范围", 10, "center"), ("等级", 20, "w"), ("含义与建议", 62, "w")]
        for c, (t, w, anch) in enumerate(heads):
            tk.Label(grid, text=t, font=FONT_B, fg=FG_DIM, bg=CARD_BG,
                     width=w, anchor=anch).grid(row=0, column=c, sticky="w", pady=(0, 6))

        cur = getattr(self, "current_iaq", None)
        cur_row = x6.iaq_guide_index(cur) if cur is not None else None

        for i, (lo, hi, name, color, advice) in enumerate(x6.IAQ_GUIDE):
            rng = "%d – %d" % (lo, hi) if hi is not None else "> %d" % lo
            mark = "◀ 当前" if i == cur_row else ""
            row_bg = "#eef4ff" if i == cur_row else CARD_BG
            tk.Label(grid, text=rng, font=FONT_MONO, fg=FG, bg=row_bg,
                     width=10, anchor="center").grid(row=i + 1, column=0, sticky="w")
            tk.Label(grid, text=name + (" " + mark if mark else ""), font=FONT_B,
                     fg=color, bg=row_bg, width=20, anchor="w").grid(
                row=i + 1, column=1, sticky="w")
            tk.Label(grid, text=advice, font=FONT, fg=FG, bg=row_bg,
                     width=62, anchor="w", wraplength=430,
                     justify="left").grid(row=i + 1, column=2, sticky="w", pady=2)

        ttk.Separator(wrap).pack(fill="x", pady=(10, 8))
        foot = ttk.Frame(wrap, style="Card.TFrame")
        foot.pack(fill="x")
        if cur is not None:
            lvl, color = x6.iaq_level(cur)
            tk.Label(foot, text="当前读数 IAQ %.1f → %s" % (cur, lvl),
                     font=FONT_B, fg=color, bg=CARD_BG).pack(side="left")
        ttk.Button(foot, text="关闭", command=win.destroy).pack(side="right")

        win.bind("<Escape>", lambda e: win.destroy())

    def _build_status(self):
        bar = ttk.Frame(self.root, padding=(10, 0))
        bar.pack(side="bottom", fill="x")
        self.status_var = tk.StringVar(value="就绪 —— 请选择串口并点击「打开串口」")
        ttk.Label(bar, textvariable=self.status_var, foreground=FG_DIM,
                  background=BG, font=FONT).pack(side="left", pady=(0, 8))
        ttk.Label(bar, text="文档：docs.waveshare.net/Environment_X6_Sensor/",
                  foreground=FG_DIM, background=BG, font=FONT).pack(side="right", pady=(0, 8))

    # ---------------------------------------------------------- 串口操作
    def refresh_ports(self, select=None):
        ports = list_serial_ports()
        self.port_map = {}
        disp = []
        for dev, desc in ports:
            d = "%s  —  %s" % (dev, desc)
            self.port_map[d] = dev
            disp.append(d)
        self.port_cb["values"] = disp
        if select:
            for k, v in self.port_map.items():
                if v == select:
                    self.port_var.set(k)
                    return
            self.port_var.set(select)
        elif disp and not self.port_var.get():
            self.port_var.set(disp[0])
        self.status_var.set("已扫描到 %d 个串口" % len(disp) if disp
                            else "未检测到串口设备，请检查 USB 连接与驱动")

    def _selected_port(self):
        v = self.port_var.get().strip()
        if not v:
            return None
        return self.port_map.get(v, v.split("  —  ")[0].strip())

    def toggle_port(self):
        if self.worker and self.worker.is_alive():
            self.close_port()
        else:
            self.open_port()

    def open_port(self):
        self._closing = False
        port = self._selected_port()
        if not port:
            messagebox.showwarning("未选择端口", "请先在下拉框里选择一个串口号。")
            return
        try:
            baud = int(self.baud_var.get())
        except ValueError:
            baud = x6.BAUDRATE

        self.worker = SerialWorker(port, baud, self.logq)
        self.worker.start()
        self.open_btn.configure(text="关闭串口")
        self.state_var.set("连接中…")
        self.state_lbl.configure(foreground=C_ACCENT)
        self.port_cb.configure(state="disabled")
        self.baud_cb.configure(state="disabled")
        self.status_var.set("正在打开 %s @ %d …" % (port, baud))

    def close_port(self):
        self._closing = True
        if self.worker:
            self.worker.stop()
            self.worker.join(timeout=3)
        self.worker = None
        self._reset_port_ui()
        self._emit("info", "串口已关闭")

    def _reset_port_ui(self):
        self.open_btn.configure(text="打开串口")
        self.state_var.set("未连接")
        self.state_lbl.configure(foreground=FG_DIM)
        self.port_cb.configure(state="normal")
        self.baud_cb.configure(state="readonly")
        self._poll_var.set(False)
        self.poll_btn.configure(text="开始连续采集")
        self._close_csv()

    # ---------------------------------------------------------- 指令发送
    def _guard(self):
        if not (self.worker and self.worker.is_alive()):
            messagebox.showwarning("未连接", "请先选择串口号并点击「打开串口」。")
            return False
        return True

    def _send(self, cmd, name, payload=b"", expect_cmd=None):
        if not self._guard():
            return
        tx, exp = x6.make_request(cmd, payload)
        self.worker.submit(name, tx, exp, expect_cmd if expect_cmd is not None else cmd)

    def _send_param(self):
        if not self._guard():
            return
        idx = {"IAQ (0x00)": 0x00, "TVOC (0x01)": 0x01,
               "HCHO (0x02)": 0x02, "CO (0x03)": 0x03}.get(self.gas_var.get(), 0x00)
        tx, exp = x6.make_request(x6.CMD_PARAM, bytes([idx]))
        self.worker.submit("读取传感器参数（%s）" % x6.GAS_ID_NAME.get(idx, idx),
                           tx, exp, x6.CMD_PARAM)

    def _send_param_all(self):
        if not self._guard():
            return
        for idx in (0x00, 0x01, 0x02, 0x03):
            tx, exp = x6.make_request(x6.CMD_PARAM, bytes([idx]))
            self.worker.submit("读取传感器参数（%s）" % x6.GAS_ID_NAME.get(idx, idx),
                               tx, exp, x6.CMD_PARAM)

    def _send_led_set(self):
        if not self._guard():
            return
        v = 0x01 if self.led_var.get().startswith("开启") else 0x00
        tx, exp = x6.make_request(x6.CMD_LED_SET, bytes([v]))
        self.worker.submit("设置运行灯（%s）" % ("开启" if v else "关闭"),
                           tx, exp, x6.CMD_LED_SET)

    def _send_raw(self):
        if not self._guard():
            return
        try:
            body = x6.parse_hex(self.raw_var.get())
        except ValueError as e:
            messagebox.showerror("十六进制有误", str(e))
            return
        if not body:
            return
        if self.raw_chk_var.get():
            tx = body                      # 用户自己带了校验和
        else:
            tx = body + bytes([x6.add8(body)])
        self.worker.submit("自定义原始帧", tx, 0, None)   # 0 = 长度未知，读到静默为止

    def toggle_poll(self):
        if not self._guard():
            self._poll_var.set(False)
            return
        if self._poll_var.get():
            self._poll_var.set(False)
            self.worker.set_poll(False, 1.0)
            self.poll_btn.configure(text="开始连续采集")
            self.status_var.set("已停止连续采集")
        else:
            try:
                itv = float(self.interval_var.get())
            except ValueError:
                itv = 1.0
            self._poll_var.set(True)
            self.worker.set_poll(True, itv)
            self.poll_btn.configure(text="停止连续采集")
            self.status_var.set("连续采集进行中：每 %.1f 秒查询一次 0x70" % itv)

    # ---------------------------------------------------------- 数据存档（xlsx，按天）
    def _toggle_csv(self):
        if self.csv_var.get():
            self._ensure_recording()
        else:
            self._close_csv()
            self.status_var.set("已停止记录数据")

    def _ensure_recording(self):
        """确保有当天的存档文件；跨天自动切换到新文件（追加到同名文件）。"""
        if not HAS_OPENPYXL:
            if not self._save_warned:
                self._save_warned = True
                self._emit("err", "缺少 openpyxl，无法记录数据，请执行：uv sync")
            return
        today = datetime.now().strftime("%Y-%m-%d")
        if self.wb is not None and self.data_date == today:
            return
        if self.wb is not None:            # 跨天：先把旧的落盘
            self._close_csv()

        try:
            os.makedirs(DATA_DIR, exist_ok=True)
            path = os.path.join(DATA_DIR, DATA_FILE_FMT % today)
            if os.path.exists(path):                       # 同一天已有 → 追加
                wb = load_workbook(path)
                ws = wb.active
                mode = "追加"
                self._build_sheet(ws, today)               # 顺带把旧表头迁移成新格式
            else:
                wb, ws = Workbook(), None
                ws = wb.active
                ws.title = "环境监测日志"
                self._build_sheet(ws, today)
                mode = "新建"
            self.wb, self.ws = wb, ws
            self.data_date = today
            self.data_path = path
            self._pending_rows = 0
            self._flush(force=True)
            self._emit("info", "开始记录数据（%s）：%s" % (mode, path))
        except Exception as e:
            self.wb = self.ws = None
            self._emit("err", "无法创建记录文件：%s" % e)

    def _build_sheet(self, ws, date_str):
        """写入/刷新标题行 + 表头行 + 列宽。对旧格式文件重复调用即可迁移。"""
        if "A1:H1" not in [str(r) for r in ws.merged_cells.ranges]:
            ws.merge_cells("A1:H1")
        t = ws.cell(row=1, column=1, value=DATA_TITLE_FMT % date_str)
        t.font = Font(bold=True, size=14)
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        head_fill = PatternFill("solid", fgColor="D9E1F2")
        for i, h in enumerate(DATA_HEADERS, start=1):
            c = ws.cell(row=2, column=i, value=h)
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.fill = head_fill
        ws.row_dimensions[2].height = 20

        for i, w in enumerate(DATA_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A3"

    def _record(self, v, lvl):
        """追加一行读数；攒够 DATA_FLUSH_ROWS 或隔 DATA_FLUSH_SECS 秒落盘一次。"""
        self._ensure_recording()
        if self.ws is None:
            return
        self.ws.append([
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            round(v["iaq"], 1), lvl,
            round(v["tvoc"], 4), round(v["hcho"], 4), round(v["co"], 4),
            round(v["temp"], 2), round(v["hum"], 2),
        ])
        r = self.ws.max_row
        for col in range(1, len(DATA_HEADERS) + 1):
            c = self.ws.cell(row=r, column=col)
            c.alignment = Alignment(horizontal="center", vertical="center")
            if col in DATA_NUM_FMT:
                c.number_format = DATA_NUM_FMT[col]
        self._pending_rows += 1
        if (self._pending_rows >= DATA_FLUSH_ROWS
                or time.time() - self._last_save >= DATA_FLUSH_SECS):
            self._flush()

    def _flush(self, force=False):
        """把内存里的工作簿写回磁盘。文件被 Excel 占用时稍后自动重试。"""
        if self.wb is None:
            return
        if not force and self._pending_rows == 0:
            return
        try:
            self.wb.save(self.data_path)
            self._pending_rows = 0
            self._last_save = time.time()
            self._save_warned = False
        except PermissionError:
            if not self._save_warned:          # 只提示一次，之后静默重试
                self._save_warned = True
                self._emit("err", "记录文件被占用（可能正用 Excel 打开），"
                                  "数据仍在内存中，关闭 Excel 后会自动写入：%s" % self.data_path)
        except Exception as e:
            self._emit("err", "写入记录文件失败：%s" % e)

    def _close_csv(self):
        self._flush(force=True)
        self.wb = None
        self.ws = None
        self.data_date = None
        self._pending_rows = 0

    # ---------------------------------------------------------- 日志
    def _emit(self, kind, title, hexs=None, lines=None):
        self.logq.put({"kind": kind, "ts": now_str(), "title": title,
                       "hex": hexs, "lines": list(lines or [])})

    def _drain_log(self):
        try:
            while True:
                item = self.logq.get_nowait()
                self._render(item)
        except queue.Empty:
            pass
        self._after_id = self.root.after(80, self._drain_log)

    def _render(self, it):
        kind = it.get("kind")
        if kind == "opened":
            self.state_var.set("已连接")
            self.state_lbl.configure(foreground=C_OK)
            self.status_var.set(it.get("title", "已连接"))
        elif kind == "closed":
            if not getattr(self, "_closing", False):   # 非主动关闭（打开失败 / 掉线）
                self._reset_port_ui()
        elif kind == "values":
            self._apply_values(it["values"])
            return
        elif kind == "status":
            self._closing = True
            self._reset_port_ui()
            self.state_var.set("连接失败")
            self.state_lbl.configure(foreground=C_ERR)
            self.status_var.set(it.get("title", ""))

        ts = it.get("ts", now_str())
        title = it.get("title", "")
        if kind == "tx":
            head, color = "▶ 发送", "tx"
        elif kind == "rx":
            head, color = "◀ 接收", "rx"
        elif kind == "err":
            head, color = "⚠ 异常", "err"
        elif kind == "summary":
            head, color = "▣ 解析", "sum"
        else:
            head, color = "ℹ 信息", "info"

        self.log.configure(state="normal")
        self.log.insert("end", "%s  " % ts, ("ts",))
        self.log.insert("end", "%s  " % head, (color,))
        self.log.insert("end", "%s\n" % title, (color,))

        if it.get("hex"):
            self.log.insert("end", "      HEX  %s\n" % it["hex"], ("mono",))
        for i, line in enumerate(it.get("lines") or []):
            branch = "└─ " if i == len(it["lines"]) - 1 else "├─ "
            self.log.insert("end", "      %s" % branch, ("dim",))
            self.log.insert("end", "%s\n" % line, (color if kind == "summary" else "mono",))

        self.log.configure(state="disabled")
        if self.autoscroll_var.get():
            self.log.see("end")

    def _apply_values(self, v):
        if "iaq" in v:
            self.val_vars["iaq"].set("%.1f" % v["iaq"])
            self.val_vars["tvoc"].set("%.3f" % v["tvoc"])
            self.val_vars["hcho"].set("%.3f" % v["hcho"])
            self.val_vars["co"].set("%.3f" % v["co"])
            self.val_vars["temp"].set("%.2f" % v["temp"])
            self.val_vars["hum"].set("%.2f" % v["hum"])
            lvl, color = x6.iaq_level(v["iaq"])
            self.level_var.set(lvl)
            self.level_lbl.configure(fg=color)
            self._level_color = color
            self.current_iaq = v["iaq"]
            if self.csv_var.get():
                self._record(v, lvl)

        extra = []
        if v.get("sn"):
            extra.append("SN = %s" % v["sn"])
        if v.get("version"):
            extra.append('固件版本 = "%s"' % v["version"])
        if extra:
            self.extra_var.set("   ｜   ".join(extra))

    def clear_log(self):
        self.log.configure(state="normal")
        self.log.delete("1.0", "end")
        self.log.configure(state="disabled")

    def save_log(self):
        path = filedialog.asksaveasfilename(
            defaultextension=".txt", filetypes=[("文本文件", "*.txt")],
            initialfile="x6_log_%s.txt" % datetime.now().strftime("%Y%m%d_%H%M%S"),
            title="保存日志")
        if not path:
            return
        try:
            with open(path, "w", encoding="utf-8") as f:
                f.write(self.log.get("1.0", "end"))
            self.status_var.set("日志已保存：%s" % path)
        except OSError as e:
            messagebox.showerror("保存失败", str(e))

    # ---------------------------------------------------------- 退出
    def on_close(self):
        if self.worker:
            self.worker.stop()
            self.worker.join(timeout=2)
        self._close_csv()
        try:
            if getattr(self, "_after_id", None):
                self.root.after_cancel(self._after_id)
        except Exception:
            pass
        self.root.destroy()


# ------------------------------------------------------------------ CLI


def offline_parse(hex_text):
    d = x6.decode(x6.parse_hex(hex_text))
    print("指令  ：%s" % d["name"])
    print("HEX   ：%s" % d["hex"])
    for label, raw, cn in d["fields"]:
        print("  %-18s %-34s %s" % (label, raw, cn))
    print("解析  ：%s" % d["summary"])


def main():
    ap = argparse.ArgumentParser(description="Environment X6 Sensor 全 API 调试工具")
    ap.add_argument("--port", help="预设串口号（仅作为默认值，仍可在界面里改）")
    ap.add_argument("--list", action="store_true", help="只列出可用串口后退出")
    ap.add_argument("--hex", help="离线解析一帧十六进制数据，例如 --hex \"70 42 30 ...\"")
    args = ap.parse_args()

    if args.list:
        for dev, desc in list_serial_ports():
            print("%-10s %s" % (dev, desc))
        return
    if args.hex:
        offline_parse(args.hex)
        return

    try:
        import tkinter  # noqa: F401
    except ImportError:
        sys.exit("当前 Python 缺少 tkinter，请在 readme 的「常见问题」里查看解决办法。")

    root = tk.Tk()
    App(root, preset_port=args.port)
    root.mainloop()


if __name__ == "__main__":
    main()
